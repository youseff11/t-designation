"""مسارات الضيوف والحضور وبوابة القاعة ولوحة متابعة العميل والتصريح.

صاحب الملف: B2. كل مسارات الميزة (views + serializers + urlpatterns) هنا
أو في ملفات جنبه بنفس البادئة — مابتلمسش ‎views.py‎ و‎serializers.py‎ (الأساس
المشترك) ولا ملفات الميزات التانية. شكل الرد والأخطاء في ‎http.py‎.
رموز القاعة والعميل والضيف هي بيانات الاعتماد: من غير تسجيل دخول، وبتتحقّق على السيرفر.

الملف ده لفريق العمل (لوحة التحكم)؛ مسارات الرموز العامة في ‎guests_public.py‎
وبتتجمّع كلها في ‎urlpatterns‎ هنا.
"""

from __future__ import annotations

import csv
import io

from django.db.models import Sum
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from django.urls import path
from django.utils import timezone
from django.utils.text import slugify
from django.views.decorators.cache import never_cache

from .. import guestexport, guestimport, qrcodes
from ..forms import GuestForm
from ..models import Guest
from ..views import _rsvp_pass_enabled
from . import guests_public
from .guests_common import (
    guest_row, invitation_basic, invitation_or_404, latest_rsvp_map, rsvp_row,
    scan_response,
)
from .http import ApiError, api_view, json_body, json_response, require_staff

XLSX_TYPE = ("application/vnd.openxmlformats-officedocument."
             "spreadsheetml.sheet")


# ---------------------------------------------------------------- كشف الضيوف
def _guests_payload(invitation) -> dict:
    guests = list(invitation.guests.all())
    rsvps = list(invitation.rsvps.all())
    latest = latest_rsvp_map(invitation)
    totals = invitation.guests.aggregate(
        allowed=Sum("entries_allowed"), used=Sum("entries_used"))
    return {
        "invitation": {
            **invitation_basic(invitation),
            "venueToken": invitation.venue_token,
            "clientToken": invitation.client_token,
        },
        "stats": {
            "guests": len(guests),
            "entriesAllowed": totals["allowed"] or 0,
            "entriesUsed": totals["used"] or 0,
            "attending": invitation.attending_count,
            "attendingPeople": invitation.total_attending_people,
            "views": invitation.public_views,
        },
        # تصريح الدخول مقفول من محرر الدعوة؟ يبقى مفيش أعمدة ولا أزرار
        # QR هنا كمان — الكشف نفسه بيفضل شغال عادي.
        "passEnabled": _rsvp_pass_enabled(invitation),
        "guests": [guest_row(g, latest.get(g.pk)) for g in guests],
        "rsvps": [rsvp_row(r) for r in rsvps],
    }


def _guest_form_data(data: dict, *, partial: Guest | None = None) -> dict:
    """حقول ‎GuestForm‎ من جسم JSON.

    الحقول الناقصة بتاخد القيمة الحالية (تعديل) أو الافتراضية (إضافة) —
    الفورم القديم كان بيبعتها كلها، والواجهة هنا ممكن تبعت جزء منها.
    """
    base = {
        "name": partial.name if partial else "",
        "phone": partial.phone if partial else "",
        "group_name": partial.group_name if partial else "",
        "plus_ones_allowed": partial.plus_ones_allowed if partial else 0,
        "note": partial.note if partial else "",
    }
    aliases = {"groupName": "group_name", "plusOnesAllowed": "plus_ones_allowed"}
    for key, value in data.items():
        key = aliases.get(key, key)
        if key in base:
            base[key] = value
    return base


def _form_errors(form) -> dict:
    return {name: " ".join(errors) for name, errors in form.errors.items()}


@never_cache
@api_view("GET", "POST")
def invitation_guests(request, pk):
    require_staff(request)
    invitation = invitation_or_404(pk)
    if request.method == "GET":
        return _guests_payload(invitation)

    form = GuestForm(_guest_form_data(json_body(request)))
    if not form.is_valid():
        raise ApiError(400, "invalid_guest", "راجع بيانات الضيف.", _form_errors(form))
    guest = form.save(commit=False)
    guest.invitation = invitation
    guest.save()
    return json_response({"guest": guest_row(guest)}, status=201)


@never_cache
@api_view("PATCH", "DELETE")
def guest_detail(request, pk):
    require_staff(request)
    guest = get_object_or_404(Guest, pk=pk)
    if request.method == "DELETE":
        guest.delete()
        return {"deleted": True}

    form = GuestForm(_guest_form_data(json_body(request), partial=guest), instance=guest)
    if not form.is_valid():
        raise ApiError(400, "invalid_guest", "راجع بيانات الضيف.", _form_errors(form))
    guest = form.save()
    return {"guest": guest_row(guest)}


@never_cache
@api_view("POST")
def guest_toggle_checkin(request, pk):
    require_staff(request)
    guest = get_object_or_404(Guest.objects.select_related("invitation"), pk=pk)
    guest.checked_in = not guest.checked_in
    guest.checked_in_at = timezone.now() if guest.checked_in else None
    guest.save(update_fields=["checked_in", "checked_in_at", "updated_at"])
    return {"checkedIn": guest.checked_in}


# ---------------------------------------------------------------- الاستيراد والتصدير
def _xlsx_to_csv(raw: bytes) -> bytes:
    """أول ورقة في ملف إكسل → CSV بنفس شكل ملف النموذج، عشان يمشي في نفس المستورد."""
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        out = io.StringIO()
        writer = csv.writer(out)
        # حد الصفوف في المستورد ‎+‎ الترويسة — مفيش داعي نقرا الباقي
        for index, row in enumerate(sheet.iter_rows(values_only=True)):
            if index > guestimport.MAX_ROWS + 1:
                break
            writer.writerow(["" if cell is None else cell for cell in row])
    finally:
        workbook.close()
    return out.getvalue().encode("utf-8")


def _is_xlsx(upload) -> bool:
    return (upload.name or "").lower().endswith(".xlsx")


@never_cache
@api_view("POST")
def guests_import(request, pk):
    require_staff(request)
    invitation = invitation_or_404(pk)
    upload = request.FILES.get("csv_file") or request.FILES.get("file")
    if upload is None:
        raise ApiError(400, "missing_file", "اختار ملف الضيوف الأول.")
    if upload.size > guestimport.MAX_BYTES:
        raise ApiError(413, "file_too_large", "الملف كبير جداً (الحد ٢ ميجا).")

    raw = upload.read()
    if _is_xlsx(upload):
        try:
            raw = _xlsx_to_csv(raw)
        except ImportError:
            raise ApiError(501, "xlsx_unavailable", "قراءة ملفات إكسل محتاجة حزمة openpyxl على الخادم.")
        except Exception:
            raise ApiError(400, "invalid_file", "ملف إكسل غير صالح.")

    report = guestimport.import_guests(invitation, raw)
    return {
        "created": report.created,
        "updated": report.updated,
        "skipped": report.skipped,
        "total": report.total,
        # نفس حد الرسائل القديم: أول ٦ أخطاء بس
        "errors": report.errors[:6],
        "errorCount": len(report.errors),
    }


@never_cache
@api_view("GET")
def guests_sample_csv(request):
    """ملف نموذجي يحمّله المستخدم ويملأه — أسهل من شرح الأعمدة بالكلام."""
    require_staff(request)
    response = HttpResponse(guestimport.SAMPLE_CSV.encode("utf-8"),
                            content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = 'attachment; filename="guests-sample.csv"'
    return response


@never_cache
@api_view("GET")
def guests_export(request, pk):
    """كشف الضيوف كملف إكسل ومعاه رموز QR — للقاعة. التوليد على السيرفر."""
    require_staff(request)
    invitation = invitation_or_404(pk)
    guests = list(invitation.guests.all())
    if request.GET.get("only", "") == "attending":
        guests = [g for g in guests if g.entries_allowed > 0]

    try:
        data = guestexport.build(
            invitation, guests,
            base_url=request.build_absolute_uri("/").rstrip("/"),
        )
    except ImportError:
        raise ApiError(501, "export_unavailable", "تصدير إكسل محتاج حزمة openpyxl على الخادم.")

    name = slugify(invitation.slug or "guests", allow_unicode=False) or "guests"
    response = HttpResponse(data, content_type=XLSX_TYPE)
    response["Content-Disposition"] = f'attachment; filename="{name}-guests.xlsx"'
    return response


# ---------------------------------------------------------------- ورقة QR والمسح
@never_cache
@api_view("GET")
def guests_qr_sheet(request, pk):
    """بيانات صفحة الطباعة: كارت لكل ضيف ومعاه الـSVG جاهز من السيرفر."""
    require_staff(request)
    invitation = invitation_or_404(pk)
    rows = []
    for guest in invitation.guests.all():
        url = request.build_absolute_uri(guest.get_absolute_url())
        rows.append({
            "id": guest.pk,
            "name": guest.name,
            "groupName": guest.group_name,
            "plusOnesAllowed": guest.plus_ones_allowed,
            "svg": qrcodes.svg_for(url, box_size=6, border=1),
        })
    return {"invitation": invitation_basic(invitation), "rows": rows}


@never_cache
@api_view("GET")
def checkin_info(request, pk):
    """شاشة الاستقبال — الأعداد الابتدائية قبل أول مسح."""
    require_staff(request)
    invitation = invitation_or_404(pk)
    totals = invitation.guests.aggregate(
        used=Sum("entries_used"), total=Sum("entries_allowed"))
    return {
        "invitation": invitation_basic(invitation),
        # بنعدّ الدخلات مش الصفوف — الضيف اللي معاه ٣ بيتحسب ٣
        "arrived": totals["used"] or 0,
        "total": totals["total"] or 0,
    }


@never_cache
@api_view("POST")
def checkin_scan(request, pk):
    """يستقبل الرمز من الماسح ويسجّل الوصول."""
    require_staff(request)
    invitation = invitation_or_404(pk)
    return scan_response(invitation, str(json_body(request).get("token") or ""))


urlpatterns = [
    path("invitations/<int:pk>/guests/", invitation_guests, name="api_guests"),
    path("invitations/<int:pk>/guests/import/", guests_import, name="api_guests_import"),
    path("invitations/<int:pk>/guests/export.xlsx", guests_export, name="api_guests_export"),
    path("invitations/<int:pk>/guests/qr-sheet/", guests_qr_sheet, name="api_guests_qr_sheet"),
    path("invitations/<int:pk>/checkin/", checkin_info, name="api_checkin"),
    path("invitations/<int:pk>/checkin/scan/", checkin_scan, name="api_checkin_scan"),
    path("guests/sample.csv", guests_sample_csv, name="api_guests_sample_csv"),
    path("guests/<int:pk>/", guest_detail, name="api_guest_detail"),
    path("guests/<int:pk>/toggle-checkin/", guest_toggle_checkin, name="api_guest_toggle_checkin"),
]
urlpatterns += guests_public.urlpatterns
