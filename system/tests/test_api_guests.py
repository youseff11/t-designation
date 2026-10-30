"""اختبارات واجهة الضيوف والحضور وبوابة القاعة والعميل والتصريح (‎/api/v1/‎).

الرموز (قاعة، عميل، ضيف) هي بيانات الاعتماد من غير تسجيل دخول، فاختبارات
الأمان هنا هي الأهم: طول الرمز، تبعيته للدعوة، حالة النشر، وCSRF.
"""

import io
import json

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client, TestCase, override_settings
from django.utils import timezone

from system.models import Customer, Guest, Invitation, RSVPResponse, Template


def _json(client, url, data=None, method="post", **extra):
    return getattr(client, method)(
        url, data=json.dumps(data if data is not None else {}),
        content_type="application/json", **extra)


# هاشر سريع للاختبارات — تشفير كلمات المرور الحقيقي بيبطّئ كل setUp
@override_settings(PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"])
class GuestsApiCase(TestCase):
    def setUp(self):
        self.template = Template.objects.create(
            name="قالب", slug="tpl", document={"version": 1, "blocks": []})
        self.customer = Customer.objects.create(name="عميل", phone="0100")
        self.inv = self._invitation("pub-one", "published")
        self.other = self._invitation("pub-two", "published")
        self.draft = self._invitation("draft-one", "draft")
        self.ali = Guest.objects.create(invitation=self.inv, name="علي", entries_allowed=2)
        self.mona = Guest.objects.create(invitation=self.inv, name="منى", entries_allowed=1)
        self.stranger = Guest.objects.create(invitation=self.other, name="غريب")
        self.draft_guest = Guest.objects.create(invitation=self.draft, name="مسودة")

        self.staff = User.objects.create_user("staff", password="Pass!12345x", is_staff=True)
        self.member = User.objects.create_user("member", password="Pass!12345x")

        self.client = Client(enforce_csrf_checks=True)

    def _invitation(self, slug, status):
        return Invitation.objects.create(
            customer=self.customer, template=self.template, slug=slug,
            name_one="ليلى", name_two="أحمد", status=status,
            event_date=timezone.now(), venue="قاعة الياسمين")

    def csrf(self):
        token = self.client.get("/api/v1/csrf/").json()["csrfToken"]
        return {"HTTP_X_CSRFTOKEN": token}

    def post(self, url, data=None, **extra):
        return _json(self.client, url, data, **self.csrf(), **extra)

    def staff_login(self):
        self.client.force_login(self.staff)


# ======================================================================
# لوحة التحكم — فريق العمل بس
# ======================================================================
class StaffPermissionTests(GuestsApiCase):
    def endpoints(self):
        pk, gid = self.inv.pk, self.ali.pk
        return [
            ("get", f"/api/v1/invitations/{pk}/guests/"),
            ("post", f"/api/v1/invitations/{pk}/guests/"),
            ("post", f"/api/v1/invitations/{pk}/guests/import/"),
            ("get", f"/api/v1/invitations/{pk}/guests/export.xlsx"),
            ("get", f"/api/v1/invitations/{pk}/guests/qr-sheet/"),
            ("get", f"/api/v1/invitations/{pk}/checkin/"),
            ("post", f"/api/v1/invitations/{pk}/checkin/scan/"),
            ("get", "/api/v1/guests/sample.csv"),
            ("patch", f"/api/v1/guests/{gid}/"),
            ("delete", f"/api/v1/guests/{gid}/"),
            ("post", f"/api/v1/guests/{gid}/toggle-checkin/"),
        ]

    def call(self, method, url):
        headers = self.csrf() if method != "get" else {}
        if method == "get":
            return self.client.get(url)
        return _json(self.client, url, {}, method=method, **headers)

    def test_anonymous_gets_401(self):
        for method, url in self.endpoints():
            with self.subTest(url=url, method=method):
                self.assertEqual(self.call(method, url).status_code, 401)

    def test_non_staff_gets_403(self):
        self.client.force_login(self.member)
        for method, url in self.endpoints():
            with self.subTest(url=url, method=method):
                self.assertEqual(self.call(method, url).status_code, 403)

    def test_nothing_changed_for_non_staff(self):
        self.client.force_login(self.member)
        self.call("delete", f"/api/v1/guests/{self.ali.pk}/")
        self.call("post", f"/api/v1/guests/{self.ali.pk}/toggle-checkin/")
        self.ali.refresh_from_db()
        self.assertFalse(self.ali.checked_in)
        self.assertTrue(Guest.objects.filter(pk=self.ali.pk).exists())

    def test_writes_require_csrf(self):
        self.staff_login()
        for method, url in [
            ("post", f"/api/v1/invitations/{self.inv.pk}/guests/"),
            ("post", f"/api/v1/guests/{self.ali.pk}/toggle-checkin/"),
            ("patch", f"/api/v1/guests/{self.ali.pk}/"),
            ("delete", f"/api/v1/guests/{self.ali.pk}/"),
            ("post", f"/api/v1/invitations/{self.inv.pk}/checkin/scan/"),
        ]:
            with self.subTest(url=url, method=method):
                self.assertEqual(
                    _json(self.client, url, {}, method=method).status_code, 403)
        self.assertTrue(Guest.objects.filter(pk=self.ali.pk).exists())


class GuestsListTests(GuestsApiCase):
    def test_payload_has_stats_guests_and_rsvps(self):
        self.staff_login()
        RSVPResponse.objects.create(
            invitation=self.inv, guest=self.ali, name="علي", status="attending",
            companions=2, message="مبروك",
            answers=[{"i": 0, "q": "نوع الأكل", "a": "لحم"}, "تالف"])
        data = self.client.get(f"/api/v1/invitations/{self.inv.pk}/guests/").json()
        self.assertEqual(data["invitation"]["id"], self.inv.pk)
        self.assertEqual(data["invitation"]["venueToken"], self.inv.venue_token)
        self.assertEqual(data["stats"]["guests"], 2)
        self.assertEqual(data["stats"]["entriesAllowed"], 3)
        self.assertEqual(data["stats"]["attending"], 1)
        self.assertEqual(data["stats"]["attendingPeople"], 3)
        self.assertTrue(data["passEnabled"])
        self.assertEqual([g["name"] for g in data["guests"]], ["علي", "منى"])
        ali = data["guests"][0]
        self.assertEqual(ali["rsvpStatus"], "attending")
        self.assertEqual(ali["passStatus"], "active")
        self.assertEqual(data["rsvps"][0]["answers"], [{"q": "نوع الأكل", "a": "لحم"}])

    def test_other_invitation_guests_are_not_listed(self):
        self.staff_login()
        data = self.client.get(f"/api/v1/invitations/{self.inv.pk}/guests/").json()
        self.assertNotIn("غريب", [g["name"] for g in data["guests"]])

    def test_unknown_invitation_is_404(self):
        self.staff_login()
        self.assertEqual(self.client.get("/api/v1/invitations/99999/guests/").status_code, 404)

    def test_pass_disabled_when_locked_from_the_editor(self):
        self.staff_login()
        self.inv.document = {"version": 1, "blocks": [
            {"id": "r", "type": "rsvp", "visible": True, "props": {"show_pass": False}}]}
        self.inv.save()
        data = self.client.get(f"/api/v1/invitations/{self.inv.pk}/guests/").json()
        self.assertFalse(data["passEnabled"])


class GuestWriteTests(GuestsApiCase):
    def test_create_guest(self):
        self.staff_login()
        response = self.post(f"/api/v1/invitations/{self.inv.pk}/guests/", {
            "name": "كريم", "phone": "0100", "groupName": "أصدقاء", "plusOnesAllowed": 2})
        self.assertEqual(response.status_code, 201)
        guest = Guest.objects.get(name="كريم")
        self.assertEqual(guest.invitation_id, self.inv.pk)
        self.assertEqual((guest.group_name, guest.plus_ones_allowed), ("أصدقاء", 2))
        # الرمز وكود التصريح بيتولّدوا على السيرفر
        self.assertEqual(len(guest.token), 24)
        self.assertTrue(guest.pass_code.startswith("FRH-"))
        self.assertEqual(response.json()["guest"]["id"], guest.pk)

    def test_create_requires_a_name(self):
        self.staff_login()
        response = self.post(f"/api/v1/invitations/{self.inv.pk}/guests/", {"name": " "})
        self.assertEqual(response.status_code, 400)
        error = response.json()["error"]
        self.assertEqual(error["code"], "invalid_guest")
        self.assertIn("name", error["fields"])
        self.assertEqual(self.inv.guests.count(), 2)

    def test_create_rejects_negative_companions(self):
        self.staff_login()
        response = self.post(f"/api/v1/invitations/{self.inv.pk}/guests/",
                             {"name": "س", "plusOnesAllowed": -1})
        self.assertEqual(response.status_code, 400)
        self.assertIn("plus_ones_allowed", response.json()["error"]["fields"])

    def test_create_rejects_non_object_body(self):
        self.staff_login()
        response = self.post(f"/api/v1/invitations/{self.inv.pk}/guests/", [1, 2])
        self.assertEqual(response.status_code, 400)

    def test_update_keeps_the_fields_not_sent(self):
        self.staff_login()
        self.ali.phone = "0111"
        self.ali.save()
        response = _json(self.client, f"/api/v1/guests/{self.ali.pk}/",
                         {"groupName": "العائلة"}, method="patch", **self.csrf())
        self.assertEqual(response.status_code, 200)
        self.ali.refresh_from_db()
        self.assertEqual((self.ali.name, self.ali.phone, self.ali.group_name),
                         ("علي", "0111", "العائلة"))
        # التعديل ما بيغيّرش الرمز ولا كود التصريح
        self.assertTrue(self.ali.token)

    def test_update_validates(self):
        self.staff_login()
        response = _json(self.client, f"/api/v1/guests/{self.ali.pk}/",
                         {"name": ""}, method="patch", **self.csrf())
        self.assertEqual(response.status_code, 400)
        self.ali.refresh_from_db()
        self.assertEqual(self.ali.name, "علي")

    def test_delete_guest(self):
        self.staff_login()
        response = _json(self.client, f"/api/v1/guests/{self.ali.pk}/",
                         method="delete", **self.csrf())
        self.assertEqual(response.status_code, 200)
        self.assertFalse(Guest.objects.filter(pk=self.ali.pk).exists())

    def test_guest_detail_unknown_is_404(self):
        self.staff_login()
        response = _json(self.client, "/api/v1/guests/99999/",
                         method="delete", **self.csrf())
        self.assertEqual(response.status_code, 404)

    def test_toggle_checkin_flips_both_ways(self):
        self.staff_login()
        url = f"/api/v1/guests/{self.ali.pk}/toggle-checkin/"
        self.assertTrue(self.post(url).json()["checkedIn"])
        self.ali.refresh_from_db()
        self.assertTrue(self.ali.checked_in)
        self.assertIsNotNone(self.ali.checked_in_at)
        self.assertFalse(self.post(url).json()["checkedIn"])
        self.ali.refresh_from_db()
        self.assertIsNone(self.ali.checked_in_at)

    def test_methods_not_allowed(self):
        self.staff_login()
        self.assertEqual(self.client.get(f"/api/v1/guests/{self.ali.pk}/").status_code, 405)
        self.assertEqual(
            self.client.get(f"/api/v1/guests/{self.ali.pk}/toggle-checkin/").status_code, 405)


class GuestImportExportTests(GuestsApiCase):
    def upload(self, name, data, content_type="text/csv"):
        self.staff_login()
        return self.client.post(
            f"/api/v1/invitations/{self.inv.pk}/guests/import/",
            {"csv_file": SimpleUploadedFile(name, data, content_type)},
            **self.csrf())

    def test_import_csv_reports_counts(self):
        csv = "الاسم,رقم الهاتف,عدد المرافقين\nسارة,0100000001,2\nعمر,,0\n".encode("utf-8")
        response = self.upload("g.csv", csv)
        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertEqual((body["created"], body["updated"], body["total"]), (2, 0, 2))
        sara = Guest.objects.get(invitation=self.inv, name="سارة")
        self.assertEqual(sara.entries_allowed, 3)
        self.assertTrue(sara.token and sara.pass_code)

    def test_import_updates_by_phone(self):
        self.ali.phone = "0100000001"
        self.ali.save()
        body = self.upload("g.csv", "name,phone\nعلي الجديد,0100000001\n".encode()).json()
        self.assertEqual((body["created"], body["updated"]), (0, 1))
        self.ali.refresh_from_db()
        self.assertEqual(self.ali.name, "علي الجديد")

    def test_import_reports_errors_capped_at_six(self):
        rows = "\n".join(f"{i}" for i in range(10))
        body = self.upload("g.csv", f"name\n{rows}\n".encode()).json()
        self.assertEqual(body["total"], 0)
        self.assertLessEqual(len(body["errors"]), 6)
        self.assertGreater(body["errorCount"], 6)

    def test_import_missing_name_column(self):
        body = self.upload("g.csv", "foo,bar\n1,2\n".encode()).json()
        self.assertEqual(body["total"], 0)
        self.assertTrue(body["errors"])

    def test_import_without_file_is_400(self):
        self.staff_login()
        response = self.client.post(
            f"/api/v1/invitations/{self.inv.pk}/guests/import/", {}, **self.csrf())
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.json()["error"]["code"], "missing_file")

    def test_import_too_large_is_rejected(self):
        from system import guestimport
        big = b"name\n" + b"x" * (guestimport.MAX_BYTES + 1)
        response = self.upload("g.csv", big)
        self.assertEqual(response.status_code, 413)
        self.assertEqual(self.inv.guests.count(), 2)

    def test_import_xlsx(self):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(["الاسم", "رقم الهاتف", "المجموعة", "عدد المرافقين"])
        ws.append(["هبة", 1000000002, "زملاء", 1])
        buf = io.BytesIO()
        wb.save(buf)
        response = self.upload("g.xlsx", buf.getvalue(),
                               "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertEqual(response.status_code, 200)
        guest = Guest.objects.get(invitation=self.inv, name="هبة")
        self.assertEqual((guest.group_name, guest.entries_allowed), ("زملاء", 2))

    def test_import_broken_xlsx_is_400(self):
        response = self.upload("g.xlsx", b"not really excel")
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.json()["error"]["code"], "invalid_file")

    def test_sample_csv(self):
        self.staff_login()
        response = self.client.get("/api/v1/guests/sample.csv")
        self.assertEqual(response.status_code, 200)
        self.assertIn("text/csv", response["Content-Type"])
        self.assertIn("attachment", response["Content-Disposition"])
        self.assertIn("الاسم", response.content.decode("utf-8"))

    def test_export_xlsx(self):
        self.staff_login()
        response = self.client.get(f"/api/v1/invitations/{self.inv.pk}/guests/export.xlsx")
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])
        self.assertIn("pub-one-guests.xlsx", response["Content-Disposition"])
        self.assertTrue(response.content.startswith(b"PK"))     # ملف zip = xlsx

    def test_export_attending_only(self):
        self.staff_login()
        self.mona.entries_allowed = 0
        self.mona.save()
        from openpyxl import load_workbook
        response = self.client.get(
            f"/api/v1/invitations/{self.inv.pk}/guests/export.xlsx?only=attending")
        sheet = load_workbook(io.BytesIO(response.content)).active
        names = [cell.value for row in sheet.iter_rows() for cell in row
                 if cell.value in ("علي", "منى")]
        self.assertEqual(names, ["علي"])


class QrSheetAndCheckinTests(GuestsApiCase):
    def test_qr_sheet_has_an_svg_per_guest(self):
        self.staff_login()
        data = self.client.get(f"/api/v1/invitations/{self.inv.pk}/guests/qr-sheet/").json()
        self.assertEqual(data["invitation"]["nameOne"], "ليلى")
        self.assertEqual(len(data["rows"]), 2)
        self.assertTrue(all(r["svg"].lstrip().startswith("<svg") for r in data["rows"]))

    def test_checkin_info_counts_entries_not_rows(self):
        self.staff_login()
        self.ali.entries_used = 1
        self.ali.save()
        data = self.client.get(f"/api/v1/invitations/{self.inv.pk}/checkin/").json()
        self.assertEqual((data["arrived"], data["total"]), (1, 3))

    def scan(self, token, invitation=None):
        invitation = invitation or self.inv
        return self.post(f"/api/v1/invitations/{invitation.pk}/checkin/scan/",
                         {"token": token})

    def test_scan_by_token_consumes_one_entry(self):
        self.staff_login()
        response = self.scan(self.ali.token)
        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertFalse(body["already"])
        self.assertEqual((body["used"], body["allowed"], body["left"]), (1, 2, 1))
        self.assertEqual((body["arrived"], body["total"]), (1, 3))
        self.assertEqual(body["name"], "علي")
        self.assertNotIn("ok", body)
        self.assertEqual(body["rsvpStatus"], "")
        self.ali.refresh_from_db()
        self.assertTrue(self.ali.checked_in)

    def test_scan_reports_the_latest_rsvp_status_code(self):
        self.staff_login()
        RSVPResponse.objects.create(invitation=self.inv, guest=self.ali, name="علي",
                                    status="maybe")
        self.assertEqual(self.scan(self.ali.token).json()["rsvpStatus"], "maybe")

    def test_scan_accepts_pass_code_and_full_link(self):
        self.staff_login()
        self.assertEqual(self.scan(self.mona.pass_code.lower()).json()["name"], "منى")
        link = f"https://x.test/i/pub-one/g/{self.ali.token}/"
        self.assertEqual(self.scan(link).json()["name"], "علي")

    def test_scan_after_the_pass_is_used_up_warns(self):
        self.staff_login()
        self.scan(self.mona.token)
        body = self.scan(self.mona.token).json()
        self.assertTrue(body["already"])
        self.assertTrue(body["error"])
        self.mona.refresh_from_db()
        self.assertEqual(self.mona.entries_used, 1)

    def test_scan_guest_without_a_pass_is_flagged(self):
        self.staff_login()
        self.mona.entries_allowed = 0
        self.mona.save()
        body = self.scan(self.mona.token).json()
        self.assertTrue(body["already"])
        self.mona.refresh_from_db()
        self.assertEqual(self.mona.entries_used, 0)

    def test_scan_token_of_another_invitation_is_404(self):
        self.staff_login()
        response = self.scan(self.stranger.token)
        self.assertEqual(response.status_code, 404)
        self.assertEqual(response.json()["error"]["code"], "guest_not_found")
        self.stranger.refresh_from_db()
        self.assertEqual(self.stranger.entries_used, 0)

    def test_scan_rejects_bad_token_length(self):
        self.staff_login()
        for token in ["", "short", "x" * 65]:
            with self.subTest(token=token):
                response = self.scan(token)
                self.assertEqual(response.status_code, 400)
                self.assertEqual(response.json()["error"]["code"], "invalid_token")


# ======================================================================
# بوابة القاعة — الرمز هو بيانات الاعتماد
# ======================================================================
class VenueApiTests(GuestsApiCase):
    def url(self, inv=None, token=None, tail=""):
        inv = inv or self.inv
        return f"/api/v1/venue/{inv.slug}/{token or inv.venue_token}/{tail}"

    def test_state_for_a_valid_token_needs_no_login(self):
        response = self.client.get(self.url())
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["X-Robots-Tag"], "noindex, nofollow")
        self.assertIn("no-store", response["Cache-Control"])
        data = response.json()
        self.assertEqual(data["invitation"]["venue"], "قاعة الياسمين")
        self.assertEqual(data["invitation"]["nameOne"], "ليلى")
        self.assertTrue(data["invitation"]["eventAtLabel"])
        self.assertEqual(data["stats"], {
            "guests": 2, "arrived": 0, "waiting": 2, "allowed": 3, "used": 0})
        self.assertEqual({r["name"] for r in data["rows"]}, {"علي", "منى"})

    def test_rows_expose_only_what_the_door_needs(self):
        row = self.client.get(self.url()).json()["rows"][0]
        self.assertEqual(set(row), {"id", "name", "code", "group", "allowed", "used",
                                    "left", "status", "checkedIn", "at"})

    def test_state_works_for_an_unpublished_invitation_like_before(self):
        self.assertEqual(self.client.get(self.url(self.draft)).status_code, 200)

    def test_wrong_token_is_404(self):
        self.assertEqual(self.client.get(self.url(token="z" * 40)).status_code, 404)

    def test_token_of_another_invitation_is_404(self):
        # رمز قاعة دعوة تانية على سلوج دعوتنا
        self.assertEqual(
            self.client.get(self.url(self.inv, token=self.other.venue_token)).status_code, 404)

    def test_client_and_guest_tokens_do_not_open_the_venue(self):
        for token in [self.inv.client_token, self.ali.token]:
            with self.subTest(token=token):
                self.assertEqual(self.client.get(self.url(token=token)).status_code, 404)

    def test_token_length_bounds(self):
        for token in ["a" * 19, "a" * 65]:
            with self.subTest(length=len(token)):
                self.assertEqual(self.client.get(self.url(token=token)).status_code, 404)

    def test_unknown_slug_is_404(self):
        response = self.client.get(f"/api/v1/venue/nope/{self.inv.venue_token}/")
        self.assertEqual(response.status_code, 404)

    def test_scan_marks_entry_and_requires_csrf(self):
        url = self.url(tail="scan/")
        blocked = _json(self.client, url, {"token": self.ali.token})
        self.assertEqual(blocked.status_code, 403)
        response = self.post(url, {"token": self.ali.token})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["used"], 1)
        self.ali.refresh_from_db()
        self.assertEqual(self.ali.entries_used, 1)

    def test_scan_cannot_reach_another_invitations_guest(self):
        response = self.post(self.url(tail="scan/"), {"token": self.stranger.token})
        self.assertEqual(response.status_code, 404)
        self.stranger.refresh_from_db()
        self.assertEqual(self.stranger.entries_used, 0)

    def test_scan_with_a_wrong_venue_token_is_404_and_changes_nothing(self):
        response = self.post(self.url(token="q" * 40, tail="scan/"), {"token": self.ali.token})
        self.assertEqual(response.status_code, 404)
        self.ali.refresh_from_db()
        self.assertEqual(self.ali.entries_used, 0)

    def test_scan_rejects_bad_token_length(self):
        self.assertEqual(self.post(self.url(tail="scan/"), {"token": "abc"}).status_code, 400)

    def mark(self, guest, action="in", **kw):
        return self.post(self.url(tail="mark/"), {"guest": guest.pk, "action": action}, **kw)

    def test_mark_in_then_undo(self):
        body = self.mark(self.ali).json()
        self.assertEqual(body["error"], "")
        self.assertEqual(body["guest"]["used"], 1)
        self.assertEqual(body["stats"]["arrived"], 1)
        body = self.mark(self.ali, "undo").json()
        self.assertEqual(body["guest"]["used"], 0)
        self.assertFalse(body["guest"]["checkedIn"])
        self.assertEqual(body["stats"]["arrived"], 0)

    def test_mark_stops_when_the_pass_is_used_up(self):
        self.mark(self.mona)
        body = self.mark(self.mona).json()
        self.assertTrue(body["error"])
        self.mona.refresh_from_db()
        self.assertEqual(self.mona.entries_used, 1)

    def test_mark_guest_without_a_pass(self):
        self.mona.entries_allowed = 0
        self.mona.save()
        self.assertTrue(self.mark(self.mona).json()["error"])
        self.mona.refresh_from_db()
        self.assertEqual(self.mona.entries_used, 0)

    def test_undo_never_goes_below_zero(self):
        self.mark(self.ali, "undo")
        self.ali.refresh_from_db()
        self.assertEqual(self.ali.entries_used, 0)

    def test_mark_guest_of_another_invitation_is_404(self):
        self.assertEqual(self.mark(self.stranger).status_code, 404)
        self.stranger.refresh_from_db()
        self.assertEqual(self.stranger.entries_used, 0)

    def test_mark_bad_guest_id_is_404(self):
        for guest in ["abc", None, 0, 99999]:
            with self.subTest(guest=guest):
                response = self.post(self.url(tail="mark/"), {"guest": guest})
                self.assertEqual(response.status_code, 404)

    def test_mark_with_a_wrong_venue_token_is_404(self):
        response = self.post(self.url(token="q" * 40, tail="mark/"),
                             {"guest": self.ali.pk})
        self.assertEqual(response.status_code, 404)

    def test_mark_requires_csrf(self):
        response = _json(self.client, self.url(tail="mark/"), {"guest": self.ali.pk})
        self.assertEqual(response.status_code, 403)


# ======================================================================
# متابعة العميل
# ======================================================================
class ClientFollowupApiTests(GuestsApiCase):
    def url(self, inv=None, token=None):
        inv = inv or self.inv
        return f"/api/v1/client/{inv.slug}/{token or inv.client_token}/"

    def test_valid_token_shows_stats_without_login(self):
        RSVPResponse.objects.create(invitation=self.inv, guest=self.ali, name="علي",
                                    status="attending", companions=1, message="ألف مبروك")
        RSVPResponse.objects.create(invitation=self.inv, name="زائر", status="declined")
        response = self.client.get(self.url())
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["X-Robots-Tag"], "noindex, nofollow")
        data = response.json()
        self.assertTrue(data["hasRsvp"] and data["hasQr"] and data["hasGuestbook"])
        self.assertEqual(data["stats"]["rsvpTotal"], 2)
        self.assertEqual(data["stats"]["attending"], 1)
        self.assertEqual(data["stats"]["declined"], 1)
        self.assertEqual(data["stats"]["companions"], 1)
        self.assertEqual(data["stats"]["guests"], 2)
        self.assertEqual(data["stats"]["messages"], 1)
        self.assertEqual(data["messageRows"][0]["message"], "ألف مبروك")
        self.assertNotIn("venueToken", data["invitation"])
        self.assertEqual({r["name"] for r in data["qrRows"]}, {"علي", "منى"})

    def test_pass_locked_from_the_editor_hides_qr(self):
        self.inv.document = {"version": 1, "blocks": [
            {"id": "r", "type": "rsvp", "visible": True, "props": {"show_pass": False}}]}
        self.inv.save()
        data = self.client.get(self.url()).json()
        self.assertFalse(data["hasQr"])
        self.assertEqual(data["qrRows"], [])
        self.assertEqual(data["stats"]["guests"], 0)

    def test_empty_invitation_has_no_sections(self):
        empty = self._invitation("empty-one", "published")
        data = self.client.get(self.url(empty)).json()
        self.assertFalse(data["hasQr"])
        self.assertFalse(data["hasRsvp"])

    def test_answers_column_only_when_there_are_answers(self):
        RSVPResponse.objects.create(invitation=self.inv, name="س", status="attending",
                                    answers=[{"i": 0, "q": "سؤال", "a": "جواب"}])
        data = self.client.get(self.url()).json()
        self.assertTrue(data["hasAnswers"])
        self.assertEqual(data["rsvpRows"][0]["answers"], [{"q": "سؤال", "a": "جواب"}])

    def test_works_for_an_unpublished_invitation_like_before(self):
        self.assertEqual(self.client.get(self.url(self.draft)).status_code, 200)

    def test_wrong_token_is_404(self):
        self.assertEqual(self.client.get(self.url(token="z" * 40)).status_code, 404)

    def test_other_invitation_token_is_404(self):
        self.assertEqual(
            self.client.get(self.url(self.inv, token=self.other.client_token)).status_code, 404)

    def test_venue_and_guest_tokens_do_not_open_it(self):
        for token in [self.inv.venue_token, self.ali.token]:
            with self.subTest(token=token):
                self.assertEqual(self.client.get(self.url(token=token)).status_code, 404)

    def test_token_length_bounds(self):
        for token in ["a" * 19, "a" * 65]:
            with self.subTest(length=len(token)):
                self.assertEqual(self.client.get(self.url(token=token)).status_code, 404)

    def test_read_only(self):
        self.assertEqual(self.post(self.url()).status_code, 405)


# ======================================================================
# تصريح الضيف
# ======================================================================
class GuestPassApiTests(GuestsApiCase):
    def url(self, inv=None, token=None):
        return f"/api/v1/guest-pass/{(inv or self.inv).slug}/{token or self.ali.token}/"

    def test_valid_link_returns_the_pass(self):
        response = self.client.get(self.url())
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["X-Robots-Tag"], "noindex, nofollow")
        data = response.json()
        self.assertEqual(data["guest"]["name"], "علي")
        self.assertEqual(data["guest"]["passCode"], self.ali.pass_code)
        self.assertEqual(
            (data["guest"]["entriesAllowed"], data["guest"]["entriesLeft"]), (2, 2))
        self.assertEqual(data["guest"]["passStatus"], "active")
        self.assertTrue(data["qrSvg"].lstrip().startswith("<svg"))
        self.assertEqual(data["pngUrl"], f"/i/pub-one/g/{self.ali.token}/qr.png")
        self.assertEqual(data["invitation"]["nameTwo"], "أحمد")

    def test_pass_status_reflects_usage(self):
        self.ali.entries_used = 2
        self.ali.save()
        self.assertEqual(self.client.get(self.url()).json()["guest"]["passStatus"], "used")

    def test_wrong_token_is_404(self):
        self.assertEqual(self.client.get(self.url(token="z" * 24)).status_code, 404)

    def test_guest_of_another_invitation_is_404(self):
        # رمز ضيف من دعوة تانية على سلوج دعوتنا
        self.assertEqual(
            self.client.get(self.url(self.inv, token=self.stranger.token)).status_code, 404)

    def test_unpublished_invitation_is_hidden_from_visitors(self):
        response = self.client.get(self.url(self.draft, self.draft_guest.token))
        self.assertEqual(response.status_code, 404)

    def test_unpublished_invitation_is_visible_to_staff(self):
        self.staff_login()
        response = self.client.get(self.url(self.draft, self.draft_guest.token))
        self.assertEqual(response.status_code, 200)

    def test_expired_invitation_is_hidden(self):
        self.inv.expires_at = timezone.now() - timezone.timedelta(days=1)
        self.inv.save()
        self.assertEqual(self.client.get(self.url()).status_code, 404)

    def test_read_only(self):
        self.assertEqual(self.post(self.url()).status_code, 405)
